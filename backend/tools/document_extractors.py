"""
document_extractors.py — Document Text Extractors (PDF & DOCX) for Project Anara.
Safely extracts raw text, tables, and pages from binary office documents.
"""

from __future__ import annotations

import logging
import re

logger = logging.getLogger("anara.document_extractors")


def _extract_text_from_docx(file_path: str) -> str:
    """Extracts text from Microsoft Word .docx files including paragraphs and tables."""
    try:
        import docx
        doc = docx.Document(file_path)
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_txt = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_txt:
                    paragraphs.append(f"| {row_txt} |")
        return "\n\n".join(paragraphs)
    except Exception as e:
        logger.warning(f"[AgentTools] Docx extract error: {e}")
        return ""


def _extract_text_from_pdf(file_path: str) -> str:
    """Extracts raw text from PDF file with multiple fallback strategies."""
    text_content = ""
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        pages_text = []
        for idx, page in enumerate(reader.pages[:30]):
            t = page.extract_text() or ""
            if t.strip():
                pages_text.append(f"[Halaman {idx+1}]\n{t}")
        text_content = "\n\n".join(pages_text)
    except Exception:
        pass

    if not text_content:
        try:
            with open(file_path, "rb") as f:
                raw = f.read()
                matches = re.findall(rb"[(](.*?)[)]\s*Tj", raw)
                if matches:
                    text_content = " ".join([m.decode("latin1", errors="ignore") for m in matches])
        except Exception:
            pass

    return text_content.strip() or "[PDF Terdeteksi: Berisi dokumen digital visual]"
